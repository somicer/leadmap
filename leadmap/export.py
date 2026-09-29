"""Excel export: daily file (places first seen that day) + cumulative file."""
import logging
from datetime import date, datetime, time, timedelta, timezone
from zoneinfo import ZoneInfo

import pandas as pd
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

log = logging.getLogger(__name__)

COLUMNS = ["name", "phone", "address", "category", "rating", "maps_url", "first_seen"]
LOCAL_COLUMNS = ["segment", "city"] + COLUMNS  # one server can hold several segments and cities


def _frame(db, tz, where: str = "", args=(), columns=COLUMNS) -> pd.DataFrame:
    rows = db.conn.execute(
        f"SELECT {', '.join(columns)}, is_mobile FROM places {where} ORDER BY first_seen", args
    ).fetchall()
    df = pd.DataFrame([dict(r) for r in rows], columns=columns + ["is_mobile"])
    if not df.empty:
        df["first_seen"] = (pd.to_datetime(df["first_seen"], utc=True)
                            .dt.tz_convert(tz).dt.strftime("%Y-%m-%d %H:%M"))
    return df


def _write(df: pd.DataFrame, path, columns=COLUMNS):
    tmp = path.with_suffix(".tmp.xlsx")
    sheets = {"موبایل": df[df["is_mobile"] == 1][columns], "همه": df[columns]}
    with pd.ExcelWriter(tmp, engine="openpyxl") as xw:
        for name, part in sheets.items():
            part.to_excel(xw, sheet_name=name, index=False)
            ws = xw.sheets[name]
            ws.sheet_view.rightToLeft = True
            ws.freeze_panes = "A2"
            for cell in ws[1]:
                cell.font = Font(bold=True)
            for i, col in enumerate(columns, 1):
                values = [str(col)] + [str(v) for v in part[col].tolist() if v is not None]
                width = min(max((len(v) for v in values), default=8) + 2, 70)
                ws.column_dimensions[get_column_letter(i)].width = width
    tmp.replace(path)  # atomic: a reader never sees a half-written file


def export_all(cfg, db, day: date | None = None):
    tz = ZoneInfo(cfg["pacing"]["timezone"])
    day = day or datetime.now(tz).date()
    out = cfg.path("exports")

    start = datetime.combine(day, time(0), tz).astimezone(timezone.utc)
    end = start + timedelta(days=1)
    daily = _frame(db, tz, "WHERE first_seen >= ? AND first_seen < ?",
                   (start.isoformat(timespec="seconds"), end.isoformat(timespec="seconds")), columns=LOCAL_COLUMNS)
    if daily.empty:
        msg = f"daily export {day}: no new places, file skipped"
        log.info(msg)
    else:
        path = out / f"leads_{day.isoformat()}.xlsx"
        _write(daily, path, columns=LOCAL_COLUMNS)
        msg = f"daily export {day}: {len(daily)} new places ({int(daily['is_mobile'].sum())} mobile) → {path.name}"
        log.info(msg)

    allp = _frame(db, tz, columns=LOCAL_COLUMNS)
    _write(allp, out / "all_leads.xlsx", columns=LOCAL_COLUMNS)
    msg += f"; all_leads.xlsx: {len(allp)} places ({int(allp['is_mobile'].sum()) if len(allp) else 0} mobile)"
    log.info(msg)
    db.event("export", msg)
    print(msg)
