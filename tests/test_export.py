from datetime import date

import openpyxl

from leadmap.db import DB
from leadmap.export import export_all


def test_export(tmp_path):
    db = DB(tmp_path / "t.db")
    base = {"address": "کرج", "category": "بنگاه", "rating": 5.0, "lat": 35.8, "lng": 51.0}
    db.upsert_place({**base, "place_id": "a", "name": "املاک الف", "phone": "09121234567", "is_mobile": True})
    db.upsert_place({**base, "place_id": "b", "name": "املاک ب", "phone": "02634567890", "is_mobile": False})
    db.conn.execute("UPDATE places SET first_seen='2020-01-01T10:00:00+00:00' WHERE place_id='b'")
    cfg = {"pacing": {"timezone": "Asia/Tehran"}}
    cfg = type("C", (dict,), {"path": lambda self, k: tmp_path})(cfg)
    from datetime import datetime
    from zoneinfo import ZoneInfo
    today = datetime.now(ZoneInfo("Asia/Tehran")).date()
    export_all(cfg, db, today)
    wb = openpyxl.load_workbook(tmp_path / f"leads_{today}.xlsx")
    assert wb.sheetnames == ["موبایل", "همه"]
    ws = wb["همه"]
    assert ws.sheet_view.rightToLeft and ws.freeze_panes == "A2" and ws["A1"].font.bold
    assert ws["A1"].value == "segment" and ws.max_row == 2 and ws["D2"].value == "09121234567"
    allw = openpyxl.load_workbook(tmp_path / "all_leads.xlsx")
    assert allw["همه"].max_row == 3 and allw["موبایل"].max_row == 2
    # a day with nothing new: daily file skipped, cumulative still written
    export_all(cfg, db, date(2019, 1, 1))
    assert not (tmp_path / "leads_2019-01-01.xlsx").exists()
