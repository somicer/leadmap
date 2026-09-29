import json
import sqlite3
import threading
from http.server import ThreadingHTTPServer

import openpyxl
import pytest
import yaml

import leadmap.service as service_mod
from leadmap.config import ROOT, Config
from leadmap.dashboard import _Failures, make_handler
from leadmap.db import DB
from leadmap.hub import HubDB, HubError, export_hub
from leadmap.hubclient import HubClient, HubUnavailable

RE = "املاک"
FRUIT = "میوه فروشی"
FRUIT_KW = ["میوه فروشی", "میوه", "میوه فروش", "تره بار"]


def place(pid, **kw):
    return {"place_id": pid, "segment": RE, "name": f"املاک {pid}", "phone": None, "is_mobile": False,
            "first_seen": "2026-09-20T10:00:00+00:00", "last_seen": "2026-09-20T10:00:00+00:00", **kw}


@pytest.fixture
def hub(tmp_path):
    h = HubDB(tmp_path / "hub.db")
    h.set_segment(RE, ["املاک", "مشاور املاک"])
    h.set_segment(FRUIT, FRUIT_KW)
    yield h
    h.close()


def worker(hub, name):
    token = hub.add_worker(name)
    return hub.worker_by_token(token)


def job_of(reply):
    j = reply["job"]
    return j and (j["city"], j["segment"])


def jobs(hub):
    return {(j["city"], j["segment"]): j for j in hub.overview()["jobs"]}


def test_segments(hub):
    assert hub.set_segment(" میوه  فروشی ", ["میوه", "میوه", " تره بار ", ""]) == ["میوه", "تره بار"]
    assert hub.queries(FRUIT) == ["میوه", "تره بار"]                 # names are whitespace-normalized
    with pytest.raises(HubError):
        hub.set_segment("x", [])
    with pytest.raises(HubError):
        hub.add_jobs(["کرج"], "ناشناخته")                          # unknown segment
    hub.add_jobs(["کرج"], FRUIT)
    with pytest.raises(HubError):
        hub.delete_segment(FRUIT)                                   # still has a queued job
    hub.set_segment("موقت", ["x"])
    hub.delete_segment("موقت")
    assert [s["name"] for s in hub.overview()["segments"]] == [RE, FRUIT]


def test_claim_order_prefer_and_exclusive(hub):
    hub.add_jobs(["شیراز", "اصفهان"], RE)
    hub.add_jobs(["اصفهان"], FRUIT)                                  # same city, other segment: separate job
    a, b = worker(hub, "a"), worker(hub, "b")
    r = hub.claim(a, "1.1.1.1", {"prefer": [["اصفهان", FRUIT]]})
    assert job_of(r) == ("اصفهان", FRUIT) and r["job"]["queries"] == FRUIT_KW  # resumes its own job
    assert job_of(hub.claim(a, "1.1.1.1", {})) == ("اصفهان", FRUIT)            # idempotent
    assert job_of(hub.claim(b, "2.2.2.2", {"prefer": [["اصفهان", FRUIT]]})) == ("شیراز", RE)  # taken → head
    c = worker(hub, "c")
    assert job_of(hub.claim(c, "3.3.3.3", {})) == ("اصفهان", RE)
    assert hub.claim(worker(hub, "d"), "4.4.4.4", {})["job"] is None          # queue empty


def test_queue_order(hub):
    hub.add_jobs(["کرج", "قم", "یزد"], RE)
    ids = {j["city"]: j["id"] for j in hub.overview()["jobs"]}
    hub.job_action(ids["یزد"], "top")
    hub.job_action(ids["کرج"], "down")
    assert [j["city"] for j in hub.overview()["jobs"]] == ["یزد", "قم", "کرج"]


def test_complete_release_pause_revoke(hub):
    hub.add_jobs(["کرج", "قم"], RE)
    a = worker(hub, "a")
    assert job_of(hub.claim(a, "ip", {})) == ("کرج", RE)
    hub.complete(a, "ip", {"city": "کرج", "segment": RE, "tiles": {"done": 9, "empty": 1}})
    j = jobs(hub)[("کرج", RE)]
    assert j["status"] == "done" and j["tiles"] == {"done": 9, "empty": 1}
    assert job_of(hub.claim(a, "ip", {})) == ("قم", RE)
    hub.job_action(jobs(hub)[("قم", RE)]["id"], "release")               # admin takes it back
    assert hub.heartbeat(a, "ip", {"city": "قم", "segment": RE})["assignment"] is None
    hub.worker_action(a["id"], "pause")
    a = hub.conn.execute("SELECT * FROM workers WHERE id=?", (a["id"],)).fetchone()
    assert hub.claim(a, "ip", {}) == {"job": None, "paused": True}
    hub.worker_action(a["id"], "resume")
    b = worker(hub, "b")
    assert job_of(hub.claim(b, "ip", {})) == ("قم", RE)
    hub.worker_action(b["id"], "revoke")                                 # revoked → job back to queue
    assert hub.worker_by_token("x") is None
    assert jobs(hub)[("قم", RE)]["status"] == "queued"
    with pytest.raises(HubError):
        hub.job_action(jobs(hub)[("کرج", RE)]["id"], "release")          # not active


def test_keywords_change_reaches_worker(hub):
    hub.add_jobs(["کرج"], FRUIT)
    a = worker(hub, "a")
    hub.claim(a, "ip", {})
    hub.set_segment(FRUIT, ["میوه", "آبمیوه"])
    assert hub.heartbeat(a, "ip", {"city": "کرج", "segment": FRUIT})["assignment"]["queries"] == ["میوه", "آبمیوه"]


def test_fail_and_requeue(hub):
    hub.add_jobs(["ناکجاآباد"], RE)
    a = worker(hub, "a")
    hub.claim(a, "ip", {})
    hub.fail(a, "ip", {"city": "ناکجاآباد", "segment": RE, "error": "Nominatim returned nothing"})
    j = hub.overview()["jobs"][0]
    assert j["status"] == "error" and "Nominatim" in j["last_error"]
    hub.job_action(j["id"], "requeue")
    assert hub.overview()["jobs"][0]["status"] == "queued"


def test_place_merge_per_segment(hub):
    a, b = worker(hub, "a"), worker(hub, "b")
    hub.heartbeat(a, "ip", {"places": [place("p1", city="کرج", last_seen="2026-09-21T00:00:00+00:00")]})
    hub.heartbeat(b, "ip", {"places": [place("p1", city="تهران", phone="09121112233", is_mobile=True,
                                             first_seen="2026-09-19T00:00:00+00:00"), {"bad": 1}]})
    r = hub.conn.execute("SELECT * FROM places").fetchone()
    assert r["phone"] == "09121112233" and r["is_mobile"] == 1
    assert r["city"] == "کرج"                                  # first city wins
    assert r["first_seen"].startswith("2026-09-19") and r["last_seen"].startswith("2026-09-21")
    hub.heartbeat(a, "ip", {"places": [place("p1")]})          # later record without phone keeps it
    assert hub.conn.execute("SELECT is_mobile FROM places").fetchone()[0] == 1
    hub.heartbeat(a, "ip", {"places": [place("p1", segment=FRUIT)]})  # same shop, other segment: own lead
    assert hub.conn.execute("SELECT COUNT(*) FROM places WHERE place_id='p1'").fetchone()[0] == 2


def test_export_hub(hub, tmp_path):
    a = worker(hub, "a")
    hub.heartbeat(a, "ip", {"places": [place("p1", city="کرج", phone="09121112233", is_mobile=True),
                                       place("p2", city="قم"), place("p3", city="قم", segment=FRUIT)]})
    cfg = Config({"pacing": {"timezone": "Asia/Tehran"}, "paths": {"exports": str(tmp_path / "ex")}})
    export_hub(cfg, hub)
    out = tmp_path / "ex" / "hub"
    wb = openpyxl.load_workbook(out / "all_leads.xlsx")
    ws = wb["همه"]
    assert (ws["A1"].value, ws["B1"].value) == ("segment", "city") and ws.max_row == 4
    assert wb["موبایل"].max_row == 2
    assert {p.name for p in out.glob("*__*.xlsx")} == {"املاک__کرج.xlsx", "املاک__قم.xlsx", "میوه_فروشی__قم.xlsx"}


def test_migrates_pre_segment_hub_db(tmp_path):
    path = tmp_path / "old.db"
    c = sqlite3.connect(path)
    c.executescript("""
        CREATE TABLE cities (name TEXT PRIMARY KEY, position REAL NOT NULL, status TEXT NOT NULL DEFAULT 'queued',
            worker_id INTEGER, assigned_at TEXT, finished_at TEXT, last_error TEXT, tiles TEXT,
            added_at TEXT NOT NULL, updated_at TEXT NOT NULL);
        CREATE TABLE places (place_id TEXT PRIMARY KEY, name TEXT, phone_raw TEXT, phone TEXT,
            is_mobile INTEGER NOT NULL DEFAULT 0, address TEXT, lat REAL, lng REAL, category TEXT, rating REAL,
            maps_url TEXT, city TEXT, tile_id TEXT, first_seen TEXT NOT NULL, last_seen TEXT NOT NULL, worker_id INTEGER);
        CREATE INDEX hub_places_city ON places(city);
        INSERT INTO cities VALUES ('اصفهان', 1, 'active', 1, 't', NULL, NULL, '{"done": 5}', 't', 't');
        INSERT INTO places(place_id, name, city, first_seen, last_seen) VALUES ('p1', 'x', 'اصفهان', 't', 't');
    """)
    c.close()
    h = HubDB(path)
    j = h.overview()["jobs"][0]
    assert (j["city"], j["segment"], j["status"], j["tiles"], j["places"]) == ("اصفهان", RE, "active", {"done": 5}, 1)
    assert h.queries(RE) == ["املاک", "مشاور املاک"]
    h.close()
    HubDB(path).close()                                          # idempotent


# --- over HTTP: dashboard handler + worker client ------------------------------------
@pytest.fixture
def server(tmp_path):
    base = yaml.safe_load((ROOT / "config.example.yaml").read_text(encoding="utf-8"))
    base["paths"] = {k: str(tmp_path / k) for k in ("db", "exports", "logs", "cache", "debug", "profile")}
    base["paths"]["hub_db"] = str(tmp_path / "hub.db")
    base["hub"] = {"server": True}
    cfg = Config(base)
    reloader = type("R", (), {"check": lambda self: None})()
    handler = make_handler(cfg, "admin", "pw", reloader, _Failures())
    httpd = ThreadingHTTPServer(("127.0.0.1", 0), handler)
    threading.Thread(target=httpd.serve_forever, daemon=True).start()
    yield cfg, f"http://127.0.0.1:{httpd.server_address[1]}"
    httpd.shutdown()


def test_client_roundtrip(server, tmp_path):
    cfg, url = server
    hub = HubDB(cfg.path("hub_db"))
    token = hub.add_worker("w1")
    hub.set_segment(FRUIT, FRUIT_KW)
    hub.add_jobs(["کرج"], FRUIT)
    hub.close()

    db = DB(tmp_path / "worker.db")
    for i in range(4):
        db.upsert_place(place(f"x{i}", city="کرج", segment=FRUIT))
    db.upsert_place(place("x0", city="کرج", segment=RE))     # same place, other segment, same last_seen
    wcfg = {"hub": {"url": url, "token": token, "sync_batch": 2, "heartbeat_s": 3600}}
    client = HubClient(wcfg, db)
    assert client.claim() == {"job": {"city": "کرج", "segment": FRUIT, "queries": FRUIT_KW}, "paused": False}
    replies = [client.heartbeat({"phase": "scraping", "city": "کرج", "segment": FRUIT}, force=True)
               for _ in range(3)]
    assert [r["accepted"] for r in replies] == [2, 2, 1]
    assert replies[0]["assignment"]["segment"] == FRUIT
    assert client.heartbeat({"city": "کرج"}) is None               # throttled, nothing left
    hub = HubDB(cfg.path("hub_db"))
    assert hub.conn.execute("SELECT COUNT(*) FROM places").fetchone()[0] == 5
    hub.close()
    client.complete("کرج", FRUIT)

    bad = HubClient({"hub": {"url": url, "token": "nope"}}, db)
    with pytest.raises(HubUnavailable, match="token rejected"):
        bad.claim()
    down = HubClient({"hub": {"url": "http://127.0.0.1:9", "token": token}}, db)
    with pytest.raises(HubUnavailable):
        down.claim()


def test_admin_api_needs_auth_and_csrf_header(server):
    import base64
    import urllib.error
    import urllib.request
    _, url = server
    auth = "Basic " + base64.b64encode(b"admin:pw").decode()

    def post(body, headers):
        req = urllib.request.Request(f"{url}/api/hub/admin", method="POST", data=json.dumps(body).encode(),
                                     headers={"Content-Type": "application/json", **headers})
        try:
            with urllib.request.urlopen(req) as r:
                return r.status, json.loads(r.read())
        except urllib.error.HTTPError as e:
            return e.code, None

    ok = {"Authorization": auth, "X-Leadmap": "1"}
    seg = {"op": "set_segment", "name": FRUIT, "keywords": "میوه فروشی، میوه\nتره بار"}
    assert post(seg, {})[0] == 401
    assert post(seg, {"Authorization": auth})[0] == 403
    assert post(seg, ok) == (200, {"keywords": ["میوه فروشی", "میوه", "تره بار"]})
    assert post({"op": "add_jobs", "names": ["یزد"], "segments": [FRUIT]}, ok) == (200, {"added": {FRUIT: ["یزد"]}})


# --- Service in hub mode (no browser, no network) -----------------------------------
class FakeHub:
    url = "fake"

    def __init__(self, cities, segment=FRUIT, queries=FRUIT_KW):
        self.queue = [{"city": c, "segment": segment, "queries": list(queries)} for c in cities]
        self.completed, self.failed = [], []

    def claim(self):
        return {"job": self.queue[0] if self.queue else None, "paused": False}

    def complete(self, city, segment):
        self.completed.append(city)
        self.queue = [j for j in self.queue if j["city"] != city]

    def fail(self, city, segment, err):
        self.failed.append(city)
        self.queue = [j for j in self.queue if j["city"] != city]

    def heartbeat(self, report, force=False):
        return {"assignment": self.queue[0] if self.queue else None, "paused": False}


def make_service(tmp_path, monkeypatch, hub):
    base = yaml.safe_load((ROOT / "config.example.yaml").read_text(encoding="utf-8"))
    base["paths"] = {k: str(tmp_path / k) for k in ("db", "exports", "logs", "cache", "debug", "profile")}
    cfg = Config(base)
    cfg["hub"] = {}
    db = DB(tmp_path / "w.db")

    def fake_init(cfg, db, city, segment):
        if city == "bad":
            raise RuntimeError("Nominatim returned nothing for 'bad'")
        db.insert_tiles([{"id": f"{segment}/{city}:0:0", "city": city, "segment": segment, "seq": 0, "row": 0,
                          "col": 0, "lat": 1, "lng": 1, "zoom": 16,
                          "min_lat": 0, "min_lng": 0, "max_lat": 2, "max_lng": 2}])
        return {"poly": None, "tiles": 1, "added": 1}

    monkeypatch.setattr(service_mod, "init_city", fake_init)
    svc = service_mod.Service(cfg, db, None, hub=hub)
    sleeps = []
    monkeypatch.setattr(svc.pacer, "sleep", lambda s, reason="": sleeps.append(reason))
    return svc, db, sleeps


def test_service_hub_flow(tmp_path, monkeypatch):
    hub = FakeHub(["bad", "کرج", "قم"])
    svc, db, sleeps = make_service(tmp_path, monkeypatch, hub)
    assert not svc.hub_ready() and hub.failed == ["bad"]         # bad city reported, skipped
    assert svc.hub_ready() and (svc.city, svc.segment) == ("کرج", FRUIT)
    assert svc.job["queries"] == FRUIT_KW
    assert db.get_state("hub_job")["city"] == "کرج"
    assert db.next_tile("کرج", FRUIT)["id"] == f"{FRUIT}/کرج:0:0"
    db.update_tile(f"{FRUIT}/کرج:0:0", status="done")
    svc.finish_job()
    assert hub.completed == ["کرج"] and svc.job is None
    assert svc.hub_ready() and svc.city == "قم"
    svc.hub_assignment = {**hub.queue[0], "queries": ["میوه"]}     # keywords edited in the panel
    assert svc.hub_ready() and svc.job["queries"] == ["میوه"]
    svc.hub_assignment = None                                     # admin released it in the panel
    hub.queue.clear()
    assert not svc.hub_ready() and svc.job is None
    assert "hub queue empty" in sleeps


def test_service_keeps_cached_job_when_hub_down(tmp_path, monkeypatch):
    class DownHub(FakeHub):
        def claim(self):
            raise HubUnavailable("down")

        def heartbeat(self, report, force=False):
            raise HubUnavailable("down")
    svc, db, sleeps = make_service(tmp_path, monkeypatch, DownHub([]))
    assert not svc.hub_ready() and sleeps == ["hub unreachable"]  # nothing cached: wait
    db.set_state("hub_job", {"city": "کرج", "segment": FRUIT, "queries": FRUIT_KW})
    assert svc.hub_ready() and (svc.city, svc.segment) == ("کرج", FRUIT)  # cached: keep scraping


def test_heartbeat_cuts_pause_short(tmp_path, monkeypatch):
    hub = FakeHub(["کرج"])
    svc, db, _ = make_service(tmp_path, monkeypatch, hub)
    assert svc.hub_ready()
    svc.phase = "pause"
    hub.queue.clear()                                   # job deleted in the panel
    with pytest.raises(service_mod.JobChanged):
        svc.hub_heartbeat(force=True)
    svc.phase = "scraping"                              # never interrupts a tile in progress
    svc.hub_heartbeat(force=True)
    svc.set_job(None)
    svc.phase = "idle (queue empty)"
    hub.heartbeat = lambda report, force=False: {"assignment": None, "paused": False, "work_available": True}
    with pytest.raises(service_mod.JobChanged):          # new job queued: claim it now
        svc.hub_heartbeat(force=True)
